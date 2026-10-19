"""Extract text from uploaded documents and split it into proposal sections.

Supports PDF (pypdf), Word .docx (python-docx, using heading styles), Markdown,
and plain text. Sections are found from Markdown headings, numbered headings
("2. Technical Approach"), short ALL-CAPS lines, and lines that match a section
title or alias in the domain profile.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from app.domain.profile import DomainProfile, slugify

MAX_FILE_BYTES = 20 * 1024 * 1024
SUPPORTED = (".pdf", ".docx", ".md", ".markdown", ".txt", ".text", ".eml", ".log")

_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")
_NUMBERED = re.compile(r"^\s*(?:section\s+)?\d{1,2}(?:\.\d{1,2})*[.)]?\s+([A-Z][^.!?]{2,80})$", re.IGNORECASE)


class UnsupportedDocumentError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise UnsupportedDocumentError(f"{filename} is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        lines: list[str] = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style.startswith("heading") or style == "title":
                level = re.findall(r"\d", style)
                lines.append("#" * (int(level[0]) if level else 1) + " " + text)
            else:
                lines.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines).strip()
    if suffix in SUPPORTED or not suffix:
        return data.decode("utf-8", errors="replace").strip()
    raise UnsupportedDocumentError(f"Unsupported file type '{suffix}'. Use one of: {', '.join(SUPPORTED)}")


def _heading(line: str, known: set[str]) -> str | None:
    stripped = line.strip()
    if not stripped or len(stripped) > 90:
        return None
    m = _MD_HEADING.match(stripped)
    if m:
        return m.group(1).strip()
    m = _NUMBERED.match(stripped)
    if m and len(m.group(1).split()) <= 8:
        return m.group(1).strip()
    bare = stripped.rstrip(":").strip()
    if slugify(bare, fallback="") in known:
        return bare
    letters = [c for c in bare if c.isalpha()]
    if letters and bare.upper() == bare and 2 <= len(bare.split()) <= 6 and len(letters) >= 6:
        return bare.title()
    return None


def split_into_sections(text: str, profile: DomainProfile) -> dict[str, str]:
    known: set[str] = set()
    for sec in profile.sections:
        known.update({sec.key, slugify(sec.title, fallback=""), *(slugify(a, fallback="") for a in sec.aliases)})
    known.discard("")

    sections: dict[str, list[str]] = {}
    current = "full_document"
    for line in (text or "").splitlines():
        title = _heading(line, known)
        if title:
            current = profile.canonical_section_key(slugify(title, fallback="section"))
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    out = {key: "\n".join(lines).strip() for key, lines in sections.items()}
    out = {key: value for key, value in out.items() if len(value) >= 40}
    if len(out) > 1 and "full_document" in out and len(out["full_document"]) < 400:
        out.pop("full_document")  # front matter before the first heading
    return out


def document_to_record(
    filename: str,
    data: bytes,
    profile: DomainProfile,
    *,
    title: str | None = None,
    client_name: str | None = None,
    industry: str | None = None,
    solution_type: str | None = None,
) -> dict:
    text = extract_text(filename, data)
    if not text:
        raise UnsupportedDocumentError(f"No text could be extracted from {filename} (scanned PDFs need OCR first)")
    stem = Path(filename).stem
    return {
        "id": f"upload-{slugify(stem, fallback='document')}",
        "title": title or stem.replace("_", " ").replace("-", " ").strip().title(),
        "client_name": client_name,
        "industry": industry,
        "solution_type": solution_type or "general",
        "proposal_sections": split_into_sections(text, profile),
    }
